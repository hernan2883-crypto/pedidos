import streamlit as st
import pandas as pd
from openpyxl import load_workbook

# 1. Configuración obligatoria al inicio
st.set_page_config(page_title="Reparto Pan", page_icon="🍞", layout="centered")

# --- CSS: DISEÑO VERTICAL OPTIMIZADO PARA LA CALLE ---
st.markdown("""
    <style>
    /* Achicar márgenes para aprovechar la pantalla del celular al máximo */
    .block-container {
        padding-top: 15px !important;
        padding-bottom: 15px !important;
        padding-left: 10px !important;
        padding-right: 10px !important;
    }
    
    /* 1. DISEÑO DEL INPUT NUMÉRICO GIGANTE */
    div[data-testid="stNumberInput"] div[data-baseweb="input"] {
        background-color: #ffffff !important;
        border: 3px solid #27AE60 !important;
        border-radius: 15px !important;
        height: 90px !important;
        box-shadow: inset 0px 3px 6px rgba(0,0,0,0.1);
    }
    div[data-testid="stNumberInput"] input {
        font-size: 45px !important;
        font-weight: 900 !important;
        text-align: center !important;
        color: #2C3E50 !important;
    }
    /* Ocultar las flechitas chiquitas de subir/bajar número */
    button[aria-label="Step Up"], button[aria-label="Step Down"] {
        display: none !important;
    }

    /* 2. BOTÓN CARGAR PAGO GIGANTE */
    div.marcador-cargar + div.element-container div.stButton > button {
        background-color: #27AE60 !important;
        color: #FFFFFF !important;
        height: 85px !important;
        font-size: 26px !important;
        font-weight: 900 !important;
        border: none !important;
        border-radius: 15px !important;
        box-shadow: 0px 6px 12px rgba(39, 174, 96, 0.4) !important;
        width: 100% !important;
        margin-top: 10px !important;
    }
    div.marcador-cargar + div.element-container div.stButton > button:active {
        background-color: #219150 !important;
        transform: scale(0.97);
    }

    /* 3. BOTÓN SALTAR CLIENTE */
    div.marcador-saltar + div.element-container div.stButton > button {
        background-color: #F8F9F9 !important;
        color: #7F8C8D !important;
        height: 60px !important;
        font-size: 18px !important;
        font-weight: bold !important;
        border: 2px solid #BDC3C7 !important;
        border-radius: 15px !important;
        width: 100% !important;
        margin-top: 15px !important;
    }
    </style>
""", unsafe_allow_html=True)

EXCEL_FILE = "Planilla_Maestra_Panaderia.xlsx"

def recargar_pantalla():
    if hasattr(st, "rerun"):
        st.rerun()
    else:
        st.experimental_rerun()

@st.cache_data(ttl=1)
def cargar_datos_completos():
    df_clientes = pd.read_excel(EXCEL_FILE, sheet_name="Clientes")
    df_control = pd.read_excel(EXCEL_FILE, sheet_name="Control_Diario")
    df_control['excel_row'] = df_control.index + 2
    df_clientes_sub = df_clientes[['ID_Cliente', 'Zona / Reparto']]
    df_unificado = df_control.merge(df_clientes_sub, on='ID_Cliente', how='left')
    return df_unificado

try:
    df_maestro = cargar_datos_completos()
except Exception as e:
    st.error(f"❌ Error al leer el Excel: {e}")
    st.stop()

if 'reparto_seleccionado' not in st.session_state:
    st.session_state.reparto_seleccionado = None
if 'cliente_actual_idx' not in st.session_state:
    st.session_state.cliente_actual_idx = 0

# --- PANTALLA 1: SELECCIÓN DE REPARTO ---
if st.session_state.reparto_seleccionado is None:
    st.markdown("<h1 style='text-align: center; color: #D35400; font-size: 36px; margin-top:20px;'>🍞 Repartos</h1>", unsafe_allow_html=True)
    st.markdown("<br>", unsafe_allow_html=True)
    
    if st.button("👨‍🍳 REPARTO P (Papá)", use_container_width=True):
        st.session_state.reparto_seleccionado = "P"
        st.session_state.cliente_actual_idx = 0
        recargar_pantalla()
    st.markdown("<br>", unsafe_allow_html=True)
    if st.button("🚚 REPARTO C (Chelo)", use_container_width=True):
        st.session_state.reparto_seleccionado = "C"
        st.session_state.cliente_actual_idx = 0
        recargar_pantalla()

# --- PANTALLA 2: SISTEMA DE COBRO RÁPIDO ---
else:
    df_reparto_del_dia = df_maestro[df_maestro['Zona / Reparto'] == st.session_state.reparto_seleccionado].copy()
    df_reparto_del_dia = df_reparto_del_dia.sort_values(by='salida').reset_index(drop=True)
    total_clientes = len(df_reparto_del_dia)
    
    if st.session_state.cliente_actual_idx >= total_clientes:
        st.balloons()
        st.success(f"🎉 ¡Reparto {st.session_state.reparto_seleccionado} Finalizado!")
        if st.button("🔄 Volver al Menú", use_container_width=True):
            st.session_state.reparto_seleccionado = None
            st.session_state.cliente_actual_idx = 0
            recargar_pantalla()
    else:
        cliente_actual = df_reparto_del_dia.iloc[st.session_state.cliente_actual_idx]
        
        # Cabecera de navegación
        col_volver, col_tit = st.columns([1, 4])
        with col_volver:
            if st.button("⬅️", use_container_width=True):
                st.session_state.reparto_seleccionado = None
                recargar_pantalla()
        with col_tit:
            st.markdown(f"<p style='text-align:right; margin-top:8px; font-weight:bold; color:#7F8C8D;'>Reparto {st.session_state.reparto_seleccionado} | Orden: #{cliente_actual['salida']}</p>", unsafe_allow_html=True)
            
        st.progress((st.session_state.cliente_actual_idx) / total_clientes)
        
        # Tarjeta del Cliente
        st.markdown(f"""
        <div style="background-color:#F8F9F9; padding:15px; border-radius:15px; border-left: 8px solid #E67E22; margin-bottom:15px; box-shadow: 0px 4px 6px rgba(0,0,0,0.05);">
            <h2 style="margin:0; color:#2C3E50; font-size:28px; line-height:1.1;">{cliente_actual['Cliente']}</h2>
            <p style="margin:5px 0 0 0; color:#95A5A6; font-size:14px;">ID: {cliente_actual['ID_Cliente']}</p>
        </div>
        """, unsafe_allow_html=True)
        
        # Deuda Anterior
        deuda = cliente_actual['Deuda Anterior']
        val_deuda = f"${deuda:,.2f}" if isinstance(deuda, (int, float)) else f"{deuda}"
        st.markdown(f"<div style='text-align:center; margin-bottom:15px;'><span style='font-size:16px; color:#7F8C8D; font-weight:bold;'>⚠️ DEUDA ANTERIOR: </span><span style='color:#C0392B; font-size:24px; font-weight:900;'>{val_deuda}</span></div>", unsafe_allow_html=True)
        
        # --- INPUT NUMÉRICO NATIVO (EL TRUCO PARA LA CALLE) ---
        st.markdown("<p style='text-align:center; font-size:18px; font-weight:bold; color:#27AE60; margin-bottom:0;'>MONTO A COBRAR:</p>", unsafe_allow_html=True)
        
        # Inicia vacío (value=None) para no tener que andar borrando el "0" antes de escribir.
        monto_ingresado = st.number_input(
            label="Monto", 
            min_value=0.0, 
            value=None, 
            step=100.0, 
            placeholder="0.00", 
            label_visibility="collapsed"
        )
        
        # --- BOTÓN CARGAR PAGO ---
        st.markdown('<div class="marcador-cargar"></div>', unsafe_allow_html=True)
        if st.button("📥 CARGAR PAGO", use_container_width=True):
            if monto_ingresado is not None and monto_ingresado > 0:
                try:
                    wb = load_workbook(EXCEL_FILE)
                    ws = wb["Control_Diario"]
                    fila_excel = int(cliente_actual['excel_row'])
                    
                    ws.cell(row=fila_excel, column=12, value=monto_ingresado)
                    wb.save(EXCEL_FILE)
                    
                    st.toast(f"✅ Guardado: ${monto_ingresado} - {cliente_actual['Cliente']}", icon="🍞")
                    st.session_state.cliente_actual_idx += 1
                    recargar_pantalla()
                except Exception as e:
                    st.error(f"Error al guardar en Excel: {e}")
            else:
                st.warning("⚠️ Ingresá un monto mayor a $0 antes de cargar.")
        
        # --- BOTÓN SALTAR ---
        st.markdown('<div class="marcador-saltar"></div>', unsafe_allow_html=True)
        if st.button("Saltar Cliente ⏭️", use_container_width=True):
            st.session_state.cliente_actual_idx += 1
            recargar_pantalla()